from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pypdf import PdfReader

ACCEPTED_EXTENSIONS = {".pdf", ".pptx", ".txt", ".md", ".csv"}


@dataclass(frozen=True)
class ParsedSection:
    text: str
    location: str


def get_file_type(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ACCEPTED_EXTENSIONS:
        accepted = ", ".join(sorted(ACCEPTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type. Accepted formats: {accepted}")
    return suffix.lstrip(".")


def parse_document(filename: str, raw_bytes: bytes) -> tuple[str, list[ParsedSection]]:
    file_type = get_file_type(filename)

    if file_type in {"txt", "md"}:
        sections = _parse_plain_text(raw_bytes)
    elif file_type == "pdf":
        sections = _parse_pdf(raw_bytes)
    elif file_type == "pptx":
        sections = _parse_pptx(raw_bytes)
    elif file_type == "csv":
        sections = _parse_csv(raw_bytes)
    else:
        raise ValueError("Unsupported file type.")

    cleaned = [section for section in sections if section.text.strip()]
    if not cleaned:
        raise ValueError("The uploaded file has no readable text.")
    return file_type, cleaned


def _decode_text(raw_bytes: bytes) -> str:
    return raw_bytes.decode("utf-8", errors="ignore").strip()


def _parse_plain_text(raw_bytes: bytes) -> list[ParsedSection]:
    return [ParsedSection(text=_decode_text(raw_bytes), location="document")]


def _parse_pdf(raw_bytes: bytes) -> list[ParsedSection]:
    try:
        reader = PdfReader(BytesIO(raw_bytes))
    except Exception as exc:
        raise ValueError("Could not read the PDF file.") from exc

    sections = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            sections.append(ParsedSection(text=text.strip(), location=f"page {index}"))
    return sections


def _parse_pptx(raw_bytes: bytes) -> list[ParsedSection]:
    try:
        presentation = Presentation(BytesIO(raw_bytes))
    except Exception as exc:
        raise ValueError("Could not read the PowerPoint file.") from exc

    sections = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())

        notes_slide = getattr(slide, "notes_slide", None)
        notes_frame = getattr(notes_slide, "notes_text_frame", None)
        if notes_frame and notes_frame.text.strip():
            parts.append(f"Speaker notes: {notes_frame.text.strip()}")

        if parts:
            sections.append(
                ParsedSection(text="\n".join(parts), location=f"slide {slide_index}")
            )
    return sections


def _parse_csv(raw_bytes: bytes) -> list[ParsedSection]:
    text = _decode_text(raw_bytes)
    try:
        dataframe = pd.read_csv(StringIO(text)).fillna("")
    except Exception as exc:
        raise ValueError("Could not read the CSV file.") from exc

    sections = []
    for index, row in dataframe.iterrows():
        values = [
            f"{column}: {value}"
            for column, value in row.items()
            if str(value).strip()
        ]
        if values:
            row_number = int(index) + 2
            sections.append(
                ParsedSection(
                    text=f"CSV row {row_number}. " + "; ".join(values),
                    location=f"row {row_number}",
                )
            )
    return sections
